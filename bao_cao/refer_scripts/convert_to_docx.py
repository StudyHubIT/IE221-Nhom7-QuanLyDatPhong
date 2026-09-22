#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
convert_to_docx.py
===================

Convert một file Markdown (mặc định: ban-hoan-chinh/Bao-cao-hoan-chinh.md —
file do refer_scripts/merge_bao_cao.py sinh ra từ nội dung trong noi-dung/)
sang file Word .docx bằng pandoc, có xử lý riêng cho ảnh SVG (Word không hiển
thị tốt ảnh SVG chèn trực tiếp từ pandoc) bằng cách render SVG -> PNG trước
khi đưa vào pandoc.

Yêu cầu cài trước:
  - pandoc (https://pandoc.org/installing.html)
  - pip install cairosvg          (để convert ảnh SVG -> PNG)
  - Google Chrome / Chromium      (khuyến nghị: sơ đồ Mermaid để chữ trong
                                    <foreignObject> HTML - cairosvg/rsvg không
                                    render được; script ưu tiên Chrome headless
                                    để giữ chữ tiếng Việt, fallback FO→<text>
                                    + cairosvg nếu không tìm thấy Chrome)
  - pip install python-docx       (dùng để sửa độ rộng cột + thêm border cho
                                    bảng, và căn giữa ảnh, sau khi pandoc
                                    convert; nếu thiếu, file vẫn ra nhưng
                                    bảng có thể lệch cột/không có viền và
                                    ảnh có thể không được căn giữa)

Cách dùng:
    python3 convert_to_docx.py
        -> convert ban-hoan-chinh/Bao-cao-hoan-chinh.md thành .docx cùng
           thư mục, dùng refer_scripts/reference-bao-cao.docx làm style (nếu
           có), và chèn refer_scripts/bia.docx làm trang bìa ở đầu tài liệu
           (nếu có).

    python3 convert_to_docx.py path/to/file.md -o path/to/output.docx
        -> convert 1 file md khác.

    python3 convert_to_docx.py --reference-doc refer_scripts/reference-bao-cao.docx
        -> chỉ định file style khác (font, size, lề trang, ...).

    python3 convert_to_docx.py --cover-doc refer_scripts/bia.docx
        -> chỉ định file .docx khác dùng làm trang bìa.

    python3 convert_to_docx.py --no-cover
        -> không chèn trang bìa.

    python3 convert_to_docx.py --keep-temp
        -> giữ lại thư mục ảnh PNG tạm để debug (mặc định script tự xoá).
"""
from __future__ import annotations

import argparse
import hashlib
import html as html_module
import os
import re
import shutil
import subprocess
import sys
import tempfile
import xml.etree.ElementTree as ET
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parent.parent
DEFAULT_MD = PROJECT_ROOT / "ban-hoan-chinh" / "Bao-cao-hoan-chinh.md"
DEFAULT_REFERENCE_DOC = Path(__file__).resolve().parent / "reference-bao-cao.docx"
DEFAULT_COVER_DOC = Path(__file__).resolve().parent / "bia.docx"

# Placeholder ký tự dùng để "che" các vùng không phải ảnh thật (code block,
# comment HTML, inline code) khi quét ảnh, độ dài giữ nguyên để offset khớp
# với text gốc.
_MASK_CHAR = "\x00"

IMAGE_RE = re.compile(r'!\[[^\]]*\]\(([^)\s]+)(?:\s+"[^"]*")?\)')

_SVG_NS = "http://www.w3.org/2000/svg"
_XLINK_NS = "http://www.w3.org/1999/xlink"
_XHTML_NS = "http://www.w3.org/1999/xhtml"
_DRAWIO_SVG_WARN = "Text is not SVG"

# Mermaid (htmlLabels:true) để nhãn trong <foreignObject> HTML. CairoSVG /
# librsvg không render foreignObject → PNG chỉ còn khung. Ưu tiên Chrome.
_CHROME_CANDIDATES = (
    "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome",
    "/Applications/Chromium.app/Contents/MacOS/Chromium",
    "/Applications/Microsoft Edge.app/Contents/MacOS/Microsoft Edge",
    "google-chrome",
    "google-chrome-stable",
    "chromium",
    "chromium-browser",
    "msedge",
)


def _svg_local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _replace_light_dark(svg_text: str) -> str:
    """Giữ màu sáng (tham số đầu) của CSS light-dark(...), kể cả rgb() lồng nhau."""
    key = "light-dark("
    out: list[str] = []
    i = 0
    lower = svg_text.lower()
    while True:
        start = lower.find(key, i)
        if start < 0:
            out.append(svg_text[i:])
            break
        out.append(svg_text[i:start])
        inner_start = start + len(key)
        depth = 1
        k = inner_start
        while k < len(svg_text) and depth:
            if svg_text[k] == "(":
                depth += 1
            elif svg_text[k] == ")":
                depth -= 1
            k += 1
        inner = svg_text[inner_start : k - 1]
        first: list[str] = []
        nest = 0
        for ch in inner:
            if ch == "(":
                nest += 1
                first.append(ch)
            elif ch == ")":
                nest -= 1
                first.append(ch)
            elif ch == "," and nest == 0:
                break
            else:
                first.append(ch)
        out.append("".join(first).strip())
        i = k
    return "".join(out)


def _find_chrome() -> str | None:
    """Trả về đường dẫn binary Chrome/Chromium/Edge nếu có trên máy."""
    for candidate in _CHROME_CANDIDATES:
        path = Path(candidate)
        if path.is_file() and os.access(path, os.X_OK):
            return str(path)
        found = shutil.which(candidate)
        if found:
            return found
    return None


def _parse_svg_pixel_size(svg_text: str) -> tuple[int, int] | None:
    """Lấy (width, height) pixel từ viewBox hoặc width/height tuyệt đối."""
    m = re.search(r'\bviewBox\s*=\s*"([^"]+)"', svg_text, re.I)
    if m:
        parts = m.group(1).replace(",", " ").split()
        if len(parts) == 4:
            try:
                w, h = float(parts[2]), float(parts[3])
            except ValueError:
                w = h = 0.0
            if w > 0 and h > 0:
                return max(1, int(round(w))), max(1, int(round(h)))

    def _abs_len(attr: str) -> float | None:
        m_attr = re.search(rf'\b{attr}\s*=\s*"([^"]+)"', svg_text, re.I)
        if not m_attr:
            return None
        raw = m_attr.group(1).strip().lower()
        if raw.endswith("%"):
            return None
        if raw.endswith("px"):
            raw = raw[:-2]
        try:
            return float(raw)
        except ValueError:
            return None

    w_abs, h_abs = _abs_len("width"), _abs_len("height")
    if w_abs and h_abs and w_abs > 0 and h_abs > 0:
        return max(1, int(round(w_abs))), max(1, int(round(h_abs)))
    return None


def _force_svg_pixel_size(svg_text: str, width: int, height: int) -> str:
    """Ép width/height pixel trên thẻ <svg> gốc (Mermaid hay để width=\"100%\")."""
    def repl_root(m: re.Match) -> str:
        tag = m.group(0)
        if re.search(r'\bwidth\s*=', tag, re.I):
            tag = re.sub(r'\bwidth\s*=\s*"[^"]*"', f'width="{width}"', tag, count=1, flags=re.I)
        else:
            tag = tag[:-1] + f' width="{width}">'
        if re.search(r'\bheight\s*=', tag, re.I):
            tag = re.sub(r'\bheight\s*=\s*"[^"]*"', f'height="{height}"', tag, count=1, flags=re.I)
        else:
            tag = tag[:-1] + f' height="{height}">'
        return tag

    return re.sub(r"<svg\b[^>]*>", repl_root, svg_text, count=1, flags=re.I)


def svg_to_png_via_chrome(svg_text: str, out_path: Path, work_dir: Path,
                          scale: float = 2.0) -> bool:
    """Render SVG (kể cả foreignObject HTML của Mermaid) bằng Chrome headless.

    Trả về True nếu ghi được PNG; False nếu không có Chrome / không đọc được
    kích thước / Chrome lỗi."""
    chrome = _find_chrome()
    if chrome is None:
        return False
    size = _parse_svg_pixel_size(svg_text)
    if size is None:
        return False
    width, height = size
    # Chrome --window-size tối thiểu ~ vài chục px; bo cao quá thấp dễ bị clamp.
    css_w, css_h = max(width, 1), max(height, 1)
    device_scale = max(1, int(round(scale)))
    sized_svg = _force_svg_pixel_size(svg_text, css_w, css_h)
    html = (
        "<!DOCTYPE html><html><head><meta charset=\"utf-8\">"
        f"<style>html,body{{margin:0;padding:0;overflow:hidden;background:#fff;"
        f"width:{css_w}px;height:{css_h}px}}svg{{display:block}}</style>"
        f"</head><body>{sized_svg}</body></html>"
    )
    work_dir.mkdir(parents=True, exist_ok=True)
    html_path = work_dir / f"{out_path.stem}-chrome.html"
    html_path.write_text(html, encoding="utf-8")
    cmd = [
        chrome,
        "--headless=new",
        "--disable-gpu",
        "--hide-scrollbars",
        "--no-first-run",
        "--no-default-browser-check",
        f"--force-device-scale-factor={device_scale}",
        f"--window-size={css_w},{css_h}",
        f"--screenshot={out_path}",
        html_path.resolve().as_uri(),
    ]
    try:
        result = subprocess.run(cmd, capture_output=True, timeout=60)
    except (OSError, subprocess.TimeoutExpired) as e:
        print(f"  [!] Chrome headless lỗi: {e}", file=sys.stderr)
        return False
    # Chrome trên macOS đôi khi exit != 0 vì lỗi CVDisplayLink vô hại nhưng
    # vẫn ghi được PNG — chấp nhận file nếu đủ lớn.
    if out_path.is_file() and out_path.stat().st_size >= 100:
        return True
    err = (result.stderr or b"").decode(errors="ignore")[-300:]
    print(
        f"  [!] Chrome không tạo được PNG ({out_path.name}, rc={result.returncode}): {err}",
        file=sys.stderr,
    )
    return False


def _css_color_to_hex(style: str | None, default: str | None = "#333333") -> str | None:
    """Lấy giá trị `color:` đầu tiên trong style inline → #rrggbb nếu parse được."""
    if not style:
        return default
    m = re.search(r"(?:^|;)\s*color\s*:\s*([^;!]+)", style, re.I)
    if not m:
        return default
    raw = m.group(1).strip()
    m_rgb = re.match(
        r"rgb\(\s*([\d.]+)\s*,\s*([\d.]+)\s*,\s*([\d.]+)\s*\)", raw, re.I
    )
    if m_rgb:
        r, g, b = (int(float(x)) for x in m_rgb.groups())
        return f"#{r:02x}{g:02x}{b:02x}"
    if re.match(r"^#[0-9a-fA-F]{3,8}$", raw) or re.match(r"^[a-zA-Z]+$", raw):
        return raw
    return default


def _lines_from_foreign_object(fo: ET.Element) -> list[str]:
    """Trích các dòng chữ từ foreignObject HTML (Mermaid: <p>/<br>)."""
    serialized = ET.tostring(fo, encoding="unicode")
    with_breaks = re.sub(r"<[^>]*\bbr\b[^>]*/?>", "\n", serialized, flags=re.I)
    plain = re.sub(r"<[^>]+>", "", with_breaks)
    plain = html_module.unescape(plain)
    return [ln.strip() for ln in plain.splitlines() if ln.strip()]


def _svg_font_size(root: ET.Element, default: float = 16.0) -> float:
    for el in root.iter():
        if _svg_local_name(el.tag) == "style":
            m = re.search(r"font-size\s*:\s*([\d.]+)px", "".join(el.itertext()) or "")
            if m:
                try:
                    return float(m.group(1))
                except ValueError:
                    pass
    return default


def _replace_foreign_objects_with_text(root: ET.Element) -> int:
    """Thay mỗi foreignObject có chữ bằng <text>/<tspan> native SVG.

    Dùng làm fallback khi không có Chrome: cairosvg đọc được <text>, không
    đọc được HTML trong foreignObject. Trả về số nhãn đã chuyển."""
    font_size = _svg_font_size(root)
    line_height = 1.5
    replaced = 0
    for parent in list(root.iter()):
        for child in list(parent):
            if _svg_local_name(child.tag) != "foreignObject":
                continue
            try:
                width = float(child.get("width") or 0)
                height = float(child.get("height") or 0)
            except ValueError:
                width = height = 0.0
            lines = _lines_from_foreign_object(child)
            idx = list(parent).index(child)
            parent.remove(child)
            if width <= 0 or height <= 0 or not lines:
                continue

            color = _css_color_to_hex(parent.get("style")) or "#333333"
            for el in child.iter():
                if _svg_local_name(el.tag) in ("div", "span", "p"):
                    found = _css_color_to_hex(el.get("style"), default=None)
                    if found:
                        color = found

            align = "center"
            for el in child.iter():
                style = el.get("style") or ""
                m = re.search(r"text-align\s*:\s*(\w+)", style, re.I)
                if m:
                    align = m.group(1).lower()
                    break
            anchor = {"center": "middle", "right": "end", "left": "start"}.get(
                align, "middle"
            )
            x = width / 2 if anchor == "middle" else (width if anchor == "end" else 0.0)
            total_h = len(lines) * font_size * line_height
            y0 = (height - total_h) / 2 + font_size * 0.85

            text_el = ET.Element(f"{{{_SVG_NS}}}text")
            text_el.set("x", f"{x:.2f}")
            text_el.set("y", f"{y0:.2f}")
            text_el.set("text-anchor", anchor)
            text_el.set("fill", color)
            text_el.set("font-size", str(font_size))
            # Arial có glyph tiếng Việt ổn định hơn Trebuchet trên macOS/cairo.
            text_el.set("font-family", "Arial, Helvetica, sans-serif")
            for i, line in enumerate(lines):
                tspan = ET.SubElement(text_el, f"{{{_SVG_NS}}}tspan")
                tspan.set("x", f"{x:.2f}")
                if i == 0:
                    tspan.set("y", f"{y0:.2f}")
                else:
                    tspan.set("dy", f"{font_size * line_height:.2f}")
                tspan.text = line
            parent.insert(idx, text_el)
            replaced += 1
    return replaced


def sanitize_drawio_svg(svg_text: str, strip_content: bool = True) -> str:
    """Chuẩn hoá SVG trước khi đưa vào cairosvg:

    - Giữ màu sáng của CSS light-dark(...)
    - Với draw.io <switch>: bỏ foreignObject + cảnh báo 'Text is not SVG',
      giữ nhánh <text> native
    - Với Mermaid (foreignObject không có <switch>): chuyển HTML labels →
      <text>/<tspan> native thay vì xoá (xoá sẽ mất hết chữ trên sơ đồ)
    - Dọn style MathJax / nhóm requiredFeatures rỗng
    """
    svg_text = _replace_light_dark(svg_text)
    if "<switch" not in svg_text and "foreignObject" not in svg_text:
        return svg_text

    ET.register_namespace("", _SVG_NS)
    ET.register_namespace("xlink", _XLINK_NS)
    ET.register_namespace("html", _XHTML_NS)
    try:
        root = ET.fromstring(svg_text)
    except ET.ParseError:
        return svg_text

    if strip_content:
        root.attrib.pop("content", None)

    parent_map = {child: parent for parent in root.iter() for child in parent}

    def _depth(elem: ET.Element) -> int:
        depth = 0
        while elem in parent_map:
            depth += 1
            elem = parent_map[elem]
        return depth

    switches = [el for el in root.iter(f"{{{_SVG_NS}}}switch")]
    switches.sort(key=_depth, reverse=True)
    for switch in switches:
        parent = parent_map.get(switch)
        if parent is None:
            continue
        keep = None
        for child in list(switch):
            if _svg_local_name(child.tag) == "foreignObject":
                continue
            if _DRAWIO_SVG_WARN in "".join(child.itertext()):
                continue
            keep = child
            break
        idx = list(parent).index(switch)
        parent.remove(switch)
        if keep is not None:
            parent.insert(idx, keep)
        parent_map = {child: parent for parent in root.iter() for child in parent}

    # Mermaid / SVG còn foreignObject ngoài <switch>: chuyển thành <text>.
    _replace_foreign_objects_with_text(root)

    for parent in list(root.iter()):
        for child in list(parent):
            local = _svg_local_name(child.tag)
            if local == "foreignObject":
                # FO rỗng / đã được thay ở trên nhưng còn sót.
                parent.remove(child)
                continue
            if local == "style" and "mjx-container" in "".join(child.itertext()):
                parent.remove(child)
                continue
            if (local == "g"
                    and child.get("requiredFeatures")
                    and len(child) == 0
                    and not (child.text or "").strip()):
                parent.remove(child)

    return ET.tostring(root, encoding="unicode", xml_declaration=True)


def _mask(pattern: str, text: str, flags=0) -> str:
    def repl(m: re.Match) -> str:
        return _MASK_CHAR * (m.end() - m.start())

    return re.sub(pattern, repl, text, flags=flags)


def find_local_svg_images(md_text: str, md_dir: Path) -> list[tuple[int, int, Path]]:
    """Trả về list (start, end, absolute_path) của phần URL trong các thẻ
    ảnh markdown ![...](...) trỏ tới file .svg cục bộ thật sự tồn tại trên
    đĩa (bỏ qua ảnh nằm trong code block / comment HTML / inline code)."""
    scan_text = md_text
    scan_text = _mask(r"```.*?```", scan_text, flags=re.DOTALL)
    scan_text = _mask(r"<!--.*?-->", scan_text, flags=re.DOTALL)
    scan_text = _mask(r"`[^`\n]*`", scan_text)

    results = []
    for m in IMAGE_RE.finditer(scan_text):
        url = m.group(1)
        if not url.lower().endswith(".svg"):
            continue
        if url.startswith(("http://", "https://", "data:")):
            continue
        abs_path = (md_dir / url).resolve()
        if abs_path.is_file():
            results.append((m.start(1), m.end(1), abs_path))
        else:
            print(f"  [!] Không tìm thấy ảnh: {url} (bỏ qua, giữ nguyên)", file=sys.stderr)
    return results


def svg_to_png(svg_path: Path, out_dir: Path, scale: float = 2.0) -> Path | None:
    """Convert 1 file SVG sang PNG.

    Thứ tự ưu tiên:
      1. Chrome headless — nếu SVG có <foreignObject> (Mermaid htmlLabels):
         render đúng chữ HTML/tiếng Việt.
      2. cairosvg trên SVG đã sanitize (FO→<text>, draw.io <switch>, light-dark).
      3. ImageMagick `convert` trên bản đã sanitize.
    """
    out_dir.mkdir(parents=True, exist_ok=True)
    digest = hashlib.sha1(str(svg_path).encode("utf-8")).hexdigest()[:10]
    out_path = out_dir / f"{svg_path.stem}-{digest}.png"

    try:
        raw_svg = svg_path.read_text(encoding="utf-8")
    except OSError as e:
        print(f"  [!] Không đọc được {svg_path.name}: {e}", file=sys.stderr)
        return None

    raw_svg = _replace_light_dark(raw_svg)
    has_html_labels = "<foreignObject" in raw_svg

    if has_html_labels:
        chrome_work = out_dir / f"{svg_path.stem}-chrome-work"
        if svg_to_png_via_chrome(raw_svg, out_path, chrome_work, scale=scale):
            print(f"    (Chrome headless — giữ chữ foreignObject)")
            return out_path
        print(
            f"  [!] {svg_path.name}: không render được bằng Chrome, "
            "fallback FO→text + cairosvg (chữ có thể lệch font).",
            file=sys.stderr,
        )

    convert_src = svg_path
    cleaned_svg = sanitize_drawio_svg(raw_svg)
    if cleaned_svg != raw_svg or has_html_labels:
        convert_src = out_dir / f"{svg_path.stem}-sanitized.svg"
        convert_src.write_text(cleaned_svg, encoding="utf-8")

    try:
        import cairosvg  # type: ignore

        cairosvg.svg2png(url=str(convert_src), write_to=str(out_path), scale=scale)
        return out_path
    except ImportError:
        pass
    except Exception as e:
        print(f"  [!] cairosvg lỗi với {svg_path.name}: {e}", file=sys.stderr)

    if shutil.which("convert"):
        try:
            subprocess.run(
                ["convert", "-background", "white", "-density", "150",
                 str(convert_src), str(out_path)],
                check=True,
                capture_output=True,
            )
            return out_path
        except subprocess.CalledProcessError as e:
            print(
                f"  [!] ImageMagick convert lỗi với {svg_path.name}: "
                f"{e.stderr.decode(errors='ignore')}",
                file=sys.stderr,
            )

    return None


# Cột đầu tiên của bảng chỉ là "#" (đánh số thứ tự) -> đổi thành "Số thứ tự"
# khi convert sang Word cho dễ hiểu hơn là ký hiệu "#". Chỉ khớp khi "#" là
# toàn bộ nội dung cột đầu tiên (ngay sau dấu | mở đầu dòng bảng).
HASH_HEADER_CELL_RE = re.compile(r'(^\s*\|\s*)#(\s*\|)', re.MULTILINE)


def rewrite_hash_table_header(md_text: str) -> str:
    return HASH_HEADER_CELL_RE.sub(lambda m: f"{m.group(1)}Số thứ tự{m.group(2)}", md_text)


# Dòng "---" (thematic break, dùng để phân tách các mục khi ghép ở
# merge_bao_cao.py) -> pandoc convert thành 1 đường kẻ ngang (VML rectangle)
# chèn giữa 2 đoạn văn. Từ khi
# reference-bao-cao.docx đã tự ngắt trang trước mỗi Heading 1
# (pageBreakBefore), đường kẻ này chỉ còn là vệt thừa nằm giữa khoảng trắng
# cuối trang trước khi ngắt trang, không cần thiết nữa -> bỏ hẳn.
THEMATIC_BREAK_RE = re.compile(r'^[ \t]*-{3,}[ \t]*$\n?', re.MULTILINE)


def strip_thematic_breaks(md_text: str) -> str:
    return THEMATIC_BREAK_RE.sub("", md_text)


def build_docx_source(md_path: Path, tmp_image_dir: Path) -> str:
    """Đọc file markdown, convert toàn bộ ảnh SVG cục bộ sang PNG, trả về
    nội dung markdown đã thay path ảnh (ảnh khác giữ nguyên, path tương đối
    vẫn resolve đúng vì pandoc được chạy với --resource-path = thư mục md)."""
    md_text = md_path.read_text(encoding="utf-8")
    md_dir = md_path.parent

    md_text = rewrite_hash_table_header(md_text)
    md_text = strip_thematic_breaks(md_text)

    svg_matches = find_local_svg_images(md_text, md_dir)
    if not svg_matches:
        return md_text

    print(f"Tìm thấy {len(svg_matches)} ảnh SVG, đang convert sang PNG...")
    replacements = []  # (start, end, new_url)
    for start, end, svg_abs in svg_matches:
        png_path = svg_to_png(svg_abs, tmp_image_dir)
        if png_path is None:
            print(f"  [!] Không convert được {svg_abs.name}, giữ nguyên SVG (Word có thể không hiện ảnh này).",
                  file=sys.stderr)
            continue
        print(f"  - {svg_abs.relative_to(PROJECT_ROOT) if svg_abs.is_relative_to(PROJECT_ROOT) else svg_abs.name} -> {png_path.name}")
        replacements.append((start, end, f"<{png_path}>"))

    # Áp dụng thay thế từ cuối văn bản lên đầu để không lệch offset.
    for start, end, new_url in sorted(replacements, key=lambda r: r[0], reverse=True):
        md_text = md_text[:start] + new_url + md_text[end:]

    return md_text


def _tblpr_insert_ordered(tblPr, new_el, later_tags, qn) -> None:
    """Chèn `new_el` vào `tblPr` ở đúng vị trí theo thứ tự bắt buộc của
    CT_TblPrBase (ISO/IEC 29500): tblStyle, tblW, jc, tblCellSpacing,
    tblInd, tblBorders, shd, tblLayout, tblCellMar, tblLook, tblCaption,
    tblDescription. `later_tags` là các thẻ (dạng "w:xxx") PHẢI đứng sau
    `new_el` - quét theo THỨ TỰ THẬT các con hiện có trong `tblPr` (không
    theo thứ tự ưu tiên cố định) để tìm phần tử thuộc `later_tags` xuất
    hiện sớm nhất và chèn ngay trước nó; nếu không có phần tử nào như vậy,
    append() ở cuối. Cách quét theo thứ tự thật (thay vì tra từng tag theo
    danh sách ưu tiên) cần thiết vì pandoc đôi khi tự sinh sẵn các phần tử
    không đúng thứ tự chuẩn (VD bảng fallback khi ảnh SVG không convert
    được có thể có tblLook đứng trước tblLayout)."""
    later = {qn(t) for t in later_tags}
    insert_before = next((el for el in tblPr if el.tag in later), None)
    if insert_before is not None:
        insert_before.addprevious(new_el)
    else:
        tblPr.append(new_el)


def _set_table_borders(tblPr, OxmlElement, qn, sz: int = 4, color: str = "000000") -> None:
    """Thêm/ghi đè <w:tblBorders> vào tblPr: viền đơn (single) cho 4 cạnh
    ngoài + đường kẻ giữa các hàng/cột (insideH/insideV), để MỌI ô trong
    bảng đều có border khi mở bằng Word (pandoc mặc định không có border
    trừ khi style "Table" trong reference-doc đã định nghĩa sẵn)."""
    borders = tblPr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        _tblpr_insert_ordered(
            tblPr, borders,
            ("w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook", "w:tblCaption", "w:tblDescription"),
            qn,
        )
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)


def _set_cell_paragraph_alignment(tc, OxmlElement, qn, align: str = "left") -> None:
    """Ép căn lề trái cho từng đoạn văn trực tiếp trong cell, không phụ thuộc
    style "Normal" của reference-doc."""
    for p in tc.findall(qn("w:p")):
        pPr = p.find(qn("w:pPr"))
        if pPr is None:
            pPr = OxmlElement("w:pPr")
            p.insert(0, pPr)
        jc = pPr.find(qn("w:jc"))
        if jc is None:
            jc = OxmlElement("w:jc")
            pPr.append(jc)
        jc.set(qn("w:val"), align)


def _set_cell_borders(tcPr, OxmlElement, qn, sz: int = 4, color: str = "000000") -> None:
    """Thêm <w:tcBorders> vào từng ô — một số renderer (đặc biệt LibreOffice/
    Word cũ) ưu tiên border khai báo ở cấp cell hơn là ở cấp bảng, nên đặt cả
    hai để chắc chắn cell nào cũng có viền dù mở bằng app nào."""
    borders = tcPr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = f"w:{edge}"
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)


def fix_table_widths(docx_path: Path) -> None:
    """pandoc xuất bảng pipe-table (không có dấu ':' căn lề) với
    <w:tblGrid/> rỗng và tblW pct=0 -> Word/LibreOffice không biết chia độ
    rộng cột, có bảng bị co lại chỉ còn thấy cột đầu. Hàm này set lại độ
    rộng cột (chia đều theo độ rộng khả dụng của trang, trừ đi phần thụt lề
    trái nếu bảng nằm lồng trong 1 bullet/list - xem `<w:tblInd>` - để bảng
    không tràn ra ngoài lề phải) cho MỌI bảng để đảm bảo hiển thị đúng, thêm
    border cho từng cell (pandoc dùng style "Table" không có sẵn viền trừ
    khi reference-doc định nghĩa), ép căn lề trái cho nội dung cell, và gỡ cờ
    "repeat as header row" nếu có để header của mỗi bảng chỉ hiện đúng 1
    lần, không lặp lại khi bảng bị ngắt trang."""
    from docx import Document
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    doc = Document(str(docx_path))
    section = doc.sections[0]
    page_usable_width_dxa = int((section.page_width - section.left_margin - section.right_margin) / 635)
    # 635 EMU = 1 dxa (twentieth of a point); Length values từ python-docx là EMU.

    fixed = 0
    for table in doc.tables:
        tbl = table._tbl
        first_tr = tbl.find(qn("w:tr"))
        if first_tr is None:
            continue
        n_cols = len(first_tr.findall(qn("w:tc")))
        if n_cols == 0:
            continue

        tblPr = tbl.find(qn("w:tblPr"))
        # Bảng lồng trong bullet/list (VD: 1 bảng nằm ngay dưới 1 dòng gạch
        # đầu dòng) được pandoc chèn tblInd = thụt lề trái theo cấp list -
        # phải trừ đi giá trị này khỏi độ rộng khả dụng, nếu không
        # tblW = full page width + tblInd sẽ tràn ra ngoài lề phải trang.
        tblInd = tblPr.find(qn("w:tblInd"))
        indent_dxa = int(tblInd.get(qn("w:w"))) if tblInd is not None else 0
        usable_width_dxa = max(page_usable_width_dxa - indent_dxa, 1)
        col_width = usable_width_dxa // n_cols

        # tblW = 100%, layout fixed để cột không bị autofit lại theo nội dung
        tblW = tblPr.find(qn("w:tblW"))
        if tblW is None:
            tblW = OxmlElement("w:tblW")
            _tblpr_insert_ordered(
                tblPr, tblW,
                ("w:jc", "w:tblCellSpacing", "w:tblInd", "w:tblBorders", "w:shd",
                 "w:tblLayout", "w:tblCellMar", "w:tblLook", "w:tblCaption", "w:tblDescription"),
                qn,
            )
        tblW.set(qn("w:type"), "dxa")
        tblW.set(qn("w:w"), str(usable_width_dxa))

        layout = tblPr.find(qn("w:tblLayout"))
        if layout is None:
            layout = OxmlElement("w:tblLayout")
            _tblpr_insert_ordered(
                tblPr, layout,
                ("w:tblCellMar", "w:tblLook", "w:tblCaption", "w:tblDescription"),
                qn,
            )
        layout.set(qn("w:type"), "fixed")

        # border cấp bảng (áp dụng mặc định cho mọi cell chưa override riêng)
        _set_table_borders(tblPr, OxmlElement, qn)

        # tblGrid
        grid = tbl.find(qn("w:tblGrid"))
        if grid is None:
            grid = OxmlElement("w:tblGrid")
            tbl.insert(list(tbl).index(tblPr) + 1, grid)
        for child in list(grid):
            grid.remove(child)
        for _ in range(n_cols):
            gc = OxmlElement("w:gridCol")
            gc.set(qn("w:w"), str(col_width))
            grid.append(gc)

        # tcW + border mỗi cell trong mỗi row (một số renderer ưu tiên tcW/
        # tcBorders hơn gridCol/tblBorders)
        for tr in tbl.findall(qn("w:tr")):
            # Đảm bảo hàng không bị đánh dấu "repeat as header row"
            # (<w:tblHeader/> trong trPr) - nếu có, Word sẽ in lại header ở
            # mỗi trang khi bảng bị ngắt trang. Gỡ bỏ để header chỉ hiện 1 lần.
            trPr = tr.find(qn("w:trPr"))
            if trPr is not None:
                tblHeader = trPr.find(qn("w:tblHeader"))
                if tblHeader is not None:
                    trPr.remove(tblHeader)
            for tc in tr.findall(qn("w:tc")):
                tcPr = tc.find(qn("w:tcPr"))
                if tcPr is None:
                    tcPr = OxmlElement("w:tcPr")
                    tc.insert(0, tcPr)
                tcW = tcPr.find(qn("w:tcW"))
                if tcW is None:
                    tcW = OxmlElement("w:tcW")
                    tcPr.append(tcW)
                tcW.set(qn("w:type"), "dxa")
                tcW.set(qn("w:w"), str(col_width))
                _set_cell_borders(tcPr, OxmlElement, qn)
                _set_cell_paragraph_alignment(tc, OxmlElement, qn, "left")
        fixed += 1

    doc.save(str(docx_path))
    print(f"  - Đã set lại độ rộng cột + border + căn lề trái cho {fixed} bảng.")


def _list_indent_dxa(paragraph_el, numbering_root, qn) -> int:
    """Trả về thụt lề trái (dxa) áp dụng cho 1 đoạn văn: ưu tiên `w:ind` khai
    báo trực tiếp trên đoạn (pPr), nếu không có thì tra ngược `numbering.xml`
    theo numId/ilvl của `w:numPr` (đoạn nằm trong bullet/numbered list kế
    thừa thụt lề từ định nghĩa list, không khai báo `w:ind` riêng). Trả về 0
    nếu không thụt lề / không xác định được (không phải lỗi - nhiều đoạn
    không nằm trong list nào)."""
    pPr = paragraph_el.find(qn("w:pPr"))
    if pPr is None:
        return 0
    ind = pPr.find(qn("w:ind"))
    if ind is not None and ind.get(qn("w:left")) is not None:
        return int(ind.get(qn("w:left")))
    if numbering_root is None:
        return 0
    numPr = pPr.find(qn("w:numPr"))
    if numPr is None:
        return 0
    numId_el = numPr.find(qn("w:numId"))
    if numId_el is None:
        return 0
    ilvl_el = numPr.find(qn("w:ilvl"))
    ilvl = ilvl_el.get(qn("w:val")) if ilvl_el is not None else "0"

    num_el = numbering_root.find(f'{qn("w:num")}[@{qn("w:numId")}="{numId_el.get(qn("w:val"))}"]')
    abs_el = num_el.find(qn("w:abstractNumId")) if num_el is not None else None
    if abs_el is None:
        return 0
    abstract_num = numbering_root.find(
        f'{qn("w:abstractNum")}[@{qn("w:abstractNumId")}="{abs_el.get(qn("w:val"))}"]'
    )
    lvl = abstract_num.find(f'{qn("w:lvl")}[@{qn("w:ilvl")}="{ilvl}"]') if abstract_num is not None else None
    lvl_pPr = lvl.find(qn("w:pPr")) if lvl is not None else None
    lvl_ind = lvl_pPr.find(qn("w:ind")) if lvl_pPr is not None else None
    if lvl_ind is None or lvl_ind.get(qn("w:left")) is None:
        return 0
    return int(lvl_ind.get(qn("w:left")))


MAX_IMAGE_HEIGHT_RATIO = 0.4  # ảnh cao tối đa = tỉ lệ này x chiều cao phần nội dung trang


def fix_image_widths(docx_path: Path) -> None:
    """Ảnh đặt ngay dưới 1 dòng bullet (VD: ảnh kiến trúc/sơ đồ minh hoạ
    trong 1 mục) được pandoc canh độ rộng gần bằng toàn bộ độ rộng khả dụng
    của trang, không trừ đi phần thụt lề của bullet - giống lỗi ở bảng (xem
    fix_table_widths: `w:tblInd`) - nên ảnh + thụt lề vượt quá lề phải,
    tràn ra ngoài trang. Đồng thời cũng giới hạn CHIỀU CAO tối đa của mọi
    ảnh = `MAX_IMAGE_HEIGHT_RATIO` x chiều cao phần nội dung của trang
    (page height trừ lề trên/dưới), để 1 ảnh không chiếm quá nhiều diện
    tích trang theo chiều dọc.

    Hàm này thu nhỏ (giữ nguyên tỉ lệ khung hình - lấy tỉ lệ thu nhỏ nhỏ
    nhất trong 2 giới hạn rộng/cao) mọi ảnh vượt quá 1 trong 2 giới hạn
    trên; giới hạn độ rộng đã trừ thụt lề thật của đoạn chứa ảnh đó (trực
    tiếp từ `w:ind`, hoặc tra ngược `numbering.xml` nếu ảnh nằm trong
    bullet/numbered list)."""
    from docx import Document
    from docx.oxml.ns import qn

    doc = Document(str(docx_path))
    section = doc.sections[0]
    page_usable_dxa = int((section.page_width - section.left_margin - section.right_margin) / 635)
    max_height_emu = int((section.page_height - section.top_margin - section.bottom_margin)
                          * MAX_IMAGE_HEIGHT_RATIO)

    try:
        numbering_root = doc.part.numbering_part.numbering_definitions._numbering
    except Exception:
        numbering_root = None

    resized = 0
    for shape in doc.inline_shapes:
        drawing = shape._inline.getparent()
        run = drawing.getparent() if drawing is not None else None
        paragraph_el = run.getparent() if run is not None else None
        if paragraph_el is None or paragraph_el.tag != qn("w:p"):
            continue
        indent_dxa = _list_indent_dxa(paragraph_el, numbering_root, qn)
        max_width_emu = max(page_usable_dxa - indent_dxa, 1) * 635

        scale = min(1.0, max_width_emu / shape.width, max_height_emu / shape.height)
        if scale < 1.0:
            shape.width = int(shape.width * scale)
            shape.height = int(shape.height * scale)
            resized += 1

    doc.save(str(docx_path))
    print(f"  - Đã thu nhỏ {resized} ảnh vượt quá độ rộng khả dụng (trừ thụt lề bullet/list) "
          f"hoặc vượt quá {int(MAX_IMAGE_HEIGHT_RATIO * 100)}% chiều cao trang.")


# Caption hình/bảng trong thân bài: "Hình 3.1: ..." / "Bảng 2.1. ..."
# (có dấu : hoặc . ngay sau số). Không khớp câu văn xuôi kiểu
# "Bảng 3.2 liệt kê..." / "Bảng 4.5 cho thấy...".
FIGURE_TABLE_CAPTION_RE = re.compile(r'^(Hình|Bảng)\s+\d+\.\d+\s*[.:]')
LIST_SECTION_TITLES = frozenset({"DANH MỤC HÌNH ẢNH", "DANH MỤC BẢNG BIỂU"})

# Style thân bài pandoc lấy từ --reference-doc. Body Text / First Paragraph /
# Compact kế thừa Normal; ép justify (w:jc=both) trên bốn style này.
BODY_PARAGRAPH_STYLE_IDS = ("Normal", "BodyText", "FirstParagraph", "Compact")
# Heading 1–9 và TOC 1–3 basedOn Normal. Nếu không tự khai báo jc, chúng
# kế thừa canh đều của thân bài — tiêu đề / dòng mục lục sẽ giãn chữ.
# Ép left để giữ canh trái cho các phần đó.
LEFT_PARAGRAPH_STYLE_IDS = (
    "Heading1", "Heading2", "Heading3", "Heading4", "Heading5",
    "Heading6", "Heading7", "Heading8", "Heading9",
    "TOC1", "TOC2", "TOC3",
)


def _set_paragraph_style_jc(style, qn, OxmlElement, align: str) -> bool:
    """Ghi w:jc lên một paragraph style. True nếu giá trị thay đổi."""
    pPr = style.find(qn("w:pPr"))
    if pPr is None:
        pPr = OxmlElement("w:pPr")
        style.append(pPr)
    jc = pPr.find(qn("w:jc"))
    if jc is None:
        jc = OxmlElement("w:jc")
        pPr.append(jc)
    if jc.get(qn("w:val")) == align:
        return False
    jc.set(qn("w:val"), align)
    return True


def set_body_alignment_justify(docx_path: Path) -> None:
    """Ép style thân bài canh đều hai bên (w:jc=both / justify).

    Đoạn văn xuôi (Normal, Body Text, First Paragraph, Compact) dùng
    justify. Heading 1–9 và dòng mục lục (TOC 1–3) giữ canh trái vì
    pandoc dựa chúng trên Normal. Title / TOC Heading / caption / ảnh
    vẫn giữ jc riêng (center) ở các bước hậu xử lý khác. Bullet/list
    được fix_list_alignment() ép left sau bước này.
    """
    from docx import Document
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    doc = Document(str(docx_path))
    justified = 0
    lefted = 0
    for style in doc.styles.element.findall(qn("w:style")):
        if style.get(qn("w:type")) != "paragraph":
            continue
        style_id = style.get(qn("w:styleId"))
        if style_id in BODY_PARAGRAPH_STYLE_IDS:
            if _set_paragraph_style_jc(style, qn, OxmlElement, "both"):
                justified += 1
        elif style_id in LEFT_PARAGRAPH_STYLE_IDS:
            if _set_paragraph_style_jc(style, qn, OxmlElement, "left"):
                lefted += 1
    doc.save(str(docx_path))
    print(
        f"  - Đã ép {justified} style thân bài canh đều (justify); "
        f"{lefted} style tiêu đề/mục lục giữ canh trái."
    )


VERBATIM_CHAR_STYLE_IDS = frozenset({"VerbatimChar", "Verbatim", "SourceCode"})


def flatten_inline_code_style(docx_path: Path) -> None:
    """Gỡ character style Verbatim Char trên các run (sinh từ `inline code`).

    Pandoc map backtick sang style Verbatim Char (Consolas, cỡ khác thân bài).
    Gỡ w:rStyle để chữ kế thừa font/cỡ của đoạn văn. Giữ nguyên fenced code
    block (paragraph style Source Code) nếu còn.
    """
    from docx import Document
    from docx.oxml.ns import qn

    doc = Document(str(docx_path))
    fixed = 0
    for r in doc.element.body.iter(qn("w:r")):
        rPr = r.find(qn("w:rPr"))
        if rPr is None:
            continue
        rStyle = rPr.find(qn("w:rStyle"))
        if rStyle is None:
            continue
        if rStyle.get(qn("w:val")) not in VERBATIM_CHAR_STYLE_IDS:
            continue
        rPr.remove(rStyle)
        fixed += 1
    doc.save(str(docx_path))
    print(f"  - Đã gỡ style Verbatim Char trên {fixed} đoạn chữ (inline code).")


# Tiêu đề giữ đậm theo paragraph style. Không gỡ w:b trên các đoạn này.
KEEP_EMPHASIS_PARAGRAPH_STYLE_IDS = frozenset(
    {
        "Heading1", "Heading2", "Heading3", "Heading4", "Heading5",
        "Heading6", "Heading7", "Heading8", "Heading9",
        "Title", "Subtitle", "TOCHeading",
    }
)
# Character style pandoc đôi khi gắn cho *...* / **...** (thường chỉ ghi w:i / w:b).
MARKDOWN_EMPHASIS_CHAR_STYLE_IDS = frozenset(
    {"Emphasis", "Strong", "SubtleEmphasis", "IntenseEmphasis"}
)
_EMPHASIS_RUN_TAGS = ("w:i", "w:iCs", "w:b", "w:bCs")


def flatten_markdown_emphasis(docx_path: Path) -> None:
    """Gỡ chữ nghiêng/đậm sinh từ markdown *...* và **...**.

    Pandoc ghi w:i / w:b (hoặc rStyle Emphasis/Strong) trên từng run.
    Gỡ các thuộc tính đó trên thân bài, bảng, tài liệu tham khảo. Heading /
    Title / TOC Heading giữ đậm theo style riêng. Trang bìa chèn sau bước
    này (prepend_cover) nên không bị ảnh hưởng.
    """
    from docx import Document
    from docx.oxml.ns import qn

    def _p_style_id(p) -> str | None:
        pPr = p.find(qn("w:pPr"))
        if pPr is None:
            return None
        pStyle = pPr.find(qn("w:pStyle"))
        if pStyle is None:
            return None
        return pStyle.get(qn("w:val"))

    doc = Document(str(docx_path))
    italic_runs = 0
    bold_runs = 0
    for r in doc.element.body.iter(qn("w:r")):
        parent = r.getparent()
        while parent is not None and parent.tag != qn("w:p"):
            parent = parent.getparent()
        if parent is not None:
            if _p_style_id(parent) in KEEP_EMPHASIS_PARAGRAPH_STYLE_IDS:
                continue
        rPr = r.find(qn("w:rPr"))
        if rPr is None:
            continue
        rStyle = rPr.find(qn("w:rStyle"))
        if rStyle is not None and rStyle.get(qn("w:val")) in MARKDOWN_EMPHASIS_CHAR_STYLE_IDS:
            rPr.remove(rStyle)
        had_italic = False
        had_bold = False
        for tag in _EMPHASIS_RUN_TAGS:
            el = rPr.find(qn(tag))
            if el is None:
                continue
            if tag in ("w:i", "w:iCs"):
                had_italic = True
            else:
                had_bold = True
            rPr.remove(el)
        if had_italic:
            italic_runs += 1
        if had_bold:
            bold_runs += 1
    doc.save(str(docx_path))
    print(
        f"  - Đã gỡ nghiêng trên {italic_runs} đoạn chữ, đậm trên {bold_runs} "
        f"đoạn chữ (markdown * / **)."
    )


def fix_toc_heading_style(docx_path: Path) -> None:
    """Ép style TOC Heading - tiêu đề "MỤC LỤC" do pandoc --toc sinh.

    Style mặc định (pandoc / reference-doc) dựa trên Heading 1 nhưng tắt đậm
    (`w:b val=0`) và tô màu xanh accent (`365F91`). Báo cáo cần tiêu đề mục
    lục canh giữa, đậm, màu đen - không đổi các dòng TOC bên dưới.
    """
    from docx import Document
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    def _ensure(parent, tag: str):
        el = parent.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            parent.append(el)
        return el

    doc = Document(str(docx_path))
    toc_style = None
    for style in doc.styles.element.findall(qn("w:style")):
        if (style.get(qn("w:type")) == "paragraph"
                and style.get(qn("w:styleId")) == "TOCHeading"):
            toc_style = style
            break
    if toc_style is None:
        print("  [!] Không tìm thấy style TOCHeading, bỏ qua chỉnh tiêu đề MỤC LỤC.")
        return

    pPr = _ensure(toc_style, "w:pPr")
    jc = _ensure(pPr, "w:jc")
    jc.set(qn("w:val"), "center")

    rPr = _ensure(toc_style, "w:rPr")
    for tag in ("w:b", "w:bCs"):
        el = rPr.find(qn(tag))
        if el is not None:
            rPr.remove(el)
        rPr.append(OxmlElement(tag))

    color = _ensure(rPr, "w:color")
    color.set(qn("w:val"), "000000")
    for attr in ("w:themeColor", "w:themeShade", "w:themeTint"):
        qattr = qn(attr)
        if color.get(qattr) is not None:
            del color.attrib[qattr]

    # Tiêu đề MỤC LỤC nằm trong khối TOC (w:sdt), không có trong
    # Document.paragraphs - duyệt toàn bộ w:p trong body để gắn jc=center.
    for p in doc.element.body.iter(qn("w:p")):
        pPr_el = p.find(qn("w:pPr"))
        if pPr_el is None:
            continue
        pStyle = pPr_el.find(qn("w:pStyle"))
        if pStyle is None or pStyle.get(qn("w:val")) != "TOCHeading":
            continue
        text = "".join(t.text or "" for t in p.iter(qn("w:t"))).strip()
        if text != "MỤC LỤC":
            continue
        jc_p = pPr_el.find(qn("w:jc"))
        if jc_p is None:
            jc_p = OxmlElement("w:jc")
            pPr_el.append(jc_p)
        jc_p.set(qn("w:val"), "center")

    doc.save(str(docx_path))
    print("  - Đã ép tiêu đề MỤC LỤC (TOC Heading): canh giữa, đậm, màu đen.")


def fix_list_alignment(docx_path: Path) -> None:
    """Đoạn văn nằm trong bullet/numbered list (`w:numPr`) mà không tự khai
    báo `w:jc` riêng sẽ kế thừa căn lề của style "Normal" (justify). Canh
    đều hai bên làm dòng ngắn nhìn xấu (khoảng cách giữa các chữ bị giãn
    bất thường). Hàm này ép bullet/list về canh trái (left). Tên bảng / hình
    ("Bảng N.i.", "Hình N.i.") được canh giữa ở center_images() - không
    xử lý ở đây để tránh canh giữa nhầm các dòng trong DANH MỤC BẢNG BIỂU.
    """
    from docx import Document
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    doc = Document(str(docx_path))
    fixed = 0
    for p in doc.paragraphs:
        pPr = p._p.find(qn("w:pPr"))
        if pPr is None:
            continue
        if pPr.find(qn("w:numPr")) is None:
            continue
        jc = pPr.find(qn("w:jc"))
        if jc is not None and jc.get(qn("w:val")) not in (None, "both"):
            continue
        if jc is None:
            jc = OxmlElement("w:jc")
            pPr.append(jc)
        jc.set(qn("w:val"), "left")
        fixed += 1

    doc.save(str(docx_path))
    print(f"  - Đã canh trái {fixed} đoạn bullet/list.")


# Style pandoc dùng cho caption ảnh/bảng. "Hình N.i" / "Bảng N.i" sau khi
# merge_bao_cao.py đánh số có thể mang style này hoặc Body Text thường.
CAPTION_STYLE_NAMES = {"Image Caption", "Caption", "Table Caption"}


def center_images(docx_path: Path) -> None:
    """Căn giữa (1) đoạn chỉ chứa ảnh, (2) caption mang style Image/Table
    Caption, và (3) đoạn thân bài đúng mẫu "Hình N.i: ..." / "Bảng N.i. ..."
    (FIGURE_TABLE_CAPTION_RE). Bỏ qua mục DANH MỤC HÌNH ẢNH / DANH MỤC BẢNG
    BIỂU - các dòng đó format_figure_list() giữ canh trái + tab-leader.
    Đoạn có ảnh lẫn text khác (icon inline) giữ nguyên căn lề."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn

    doc = Document(str(docx_path))
    centered_images = 0
    centered_captions = 0

    def _center_image_or_styled_caption(p) -> bool:
        """True nếu đã xử lý đoạn này (caption style hoặc ảnh đứng riêng)."""
        nonlocal centered_images, centered_captions
        if p.style is not None and p.style.name in CAPTION_STYLE_NAMES:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            centered_captions += 1
            return True
        drawings = p._p.findall(".//" + qn("w:drawing"))
        if not drawings:
            return False
        text_content = "".join(t.text or "" for t in p._p.iter(qn("w:t")))
        if text_content.strip():
            return False
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        centered_images += 1
        return True

    in_list_section = False
    for p in doc.paragraphs:
        style_name = p.style.name if p.style is not None else ""
        text = p.text.strip()
        if style_name.startswith("Heading") and text in LIST_SECTION_TITLES:
            in_list_section = True
            continue
        if in_list_section and style_name.startswith("Heading"):
            in_list_section = False

        if _center_image_or_styled_caption(p):
            continue
        # Caption dạng Body Text "Hình 3.1: ..." trong thân bài - không đụng
        # danh mục (cần giữ trái cho tab-leader).
        if in_list_section:
            continue
        if FIGURE_TABLE_CAPTION_RE.match(text):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            centered_captions += 1

    # Ảnh + caption trong ô bảng (nếu có).
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if _center_image_or_styled_caption(p):
                        continue
                    if FIGURE_TABLE_CAPTION_RE.match(p.text.strip()):
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        centered_captions += 1

    doc.save(str(docx_path))
    print(f"  - Đã căn giữa {centered_images} ảnh và {centered_captions} caption.")

FIG_ANCHOR_RE = re.compile(r'^fig-\d+-\d+$')
TBL_ANCHOR_RE = re.compile(r'^tbl-\d+-\d+$')


def _add_field(paragraph, instr_text: str, placeholder: str, OxmlElement, qn) -> None:
    """Chèn 1 field Word (`begin` -> `instrText` -> `separate` -> kết quả
    placeholder -> `end`) vào cuối `paragraph` - tự cập nhật khi bấm Update
    Field (F9) hoặc khi in, giống hệt cách Table of Contents chuẩn hoạt
    động. `placeholder` chỉ là giá trị hiển thị TRƯỚC khi field được cập
    nhật lần đầu (Word tự thay bằng giá trị thật, không cần sửa gì thêm)."""
    r = paragraph.add_run()
    fld = OxmlElement("w:fldChar")
    fld.set(qn("w:fldCharType"), "begin")
    r._r.append(fld)

    r = paragraph.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instr_text
    r._r.append(instr)

    r = paragraph.add_run()
    fld = OxmlElement("w:fldChar")
    fld.set(qn("w:fldCharType"), "separate")
    r._r.append(fld)

    paragraph.add_run(placeholder)

    r = paragraph.add_run()
    fld = OxmlElement("w:fldChar")
    fld.set(qn("w:fldCharType"), "end")
    r._r.append(fld)


def _add_pageref_field(paragraph, anchor: str, OxmlElement, qn) -> None:
    """Chèn field Word `PAGEREF <anchor> \\h` - hiện số trang thật của
    bookmark `anchor`."""
    _add_field(paragraph, f" PAGEREF {anchor} \\h ", "?", OxmlElement, qn)


def _strip_hyperlink_color(hyperlink_el, OxmlElement, qn) -> None:
    """Ép màu chữ đen + bỏ gạch chân cho text trong `hyperlink_el` (đè lên
    style "Hyperlink" mặc định của Word - xanh + gạch chân), nhưng vẫn giữ
    nguyên `w:anchor` nên click vẫn nhảy tới đúng ảnh như thường."""
    for r in hyperlink_el.findall(qn("w:r")):
        rPr = r.find(qn("w:rPr"))
        if rPr is None:
            rPr = OxmlElement("w:rPr")
            r.insert(0, rPr)
        for tag in ("w:color", "w:u"):
            el = rPr.find(qn(tag))
            if el is not None:
                rPr.remove(el)
        color = OxmlElement("w:color")
        color.set(qn("w:val"), "000000")
        rPr.append(color)
        underline = OxmlElement("w:u")
        underline.set(qn("w:val"), "none")
        rPr.append(underline)


# (heading text, regex cho anchor) - mỗi mục "danh mục" tự sinh bởi
# merge_bao_cao.py (number_figures()/number_tables()) theo cùng 1 khuôn:
# mỗi dòng là 1 đoạn văn CHỈ chứa đúng 1 hyperlink nội bộ [Nhãn](#anchor).
LIST_SECTIONS = (
    ("DANH MỤC HÌNH ẢNH", FIG_ANCHOR_RE),
    ("DANH MỤC BẢNG BIỂU", TBL_ANCHOR_RE),
)

# Caption bảng còn sót `{#tbl-N-i}` (format cũ - pandoc không tạo bookmark
# từ attribute trên caption, chỉ giữ thành text). Gỡ text + đảm bảo bookmark.
LITERAL_TBL_ID_RE = re.compile(r'\s*\{#(tbl-\d+-\d+)\}\s*')
BANG_CAPTION_NUM_RE = re.compile(r'^Bảng\s+(\d+)\.(\d+)[.:]')


def _next_bookmark_id(doc) -> int:
    from docx.oxml.ns import qn
    max_id = 0
    for el in doc.element.body.iter(qn("w:bookmarkStart")):
        try:
            max_id = max(max_id, int(el.get(qn("w:id")) or 0))
        except ValueError:
            pass
    return max_id + 1


def _ensure_bookmark_on_paragraph(paragraph, name: str, bookmark_id: int, OxmlElement, qn) -> None:
    """Gắn bookmarkStart/End quanh nội dung đoạn văn nếu chưa có bookmark `name`."""
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def ensure_table_bookmarks(docx_path: Path) -> None:
    """Đảm bảo mọi anchor `tbl-N-i` dùng bởi DANH MỤC BẢNG BIỂU tồn tại dưới
    dạng bookmark Word.

    Bối cảnh: pandoc tạo bookmark từ `![](){#fig-...}` và `[]{#tbl-...}`,
    nhưng KHÔNG từ `{#tbl-...}` gắn vào dòng caption `: ...` (chuỗi đó thành
    text trong caption). merge_bao_cao.py đã chuyển sang empty span; hàm này
    là lớp phòng thủ khi convert file md cũ hoặc khi bookmark bị thiếu:
      1) Gỡ literal `{#tbl-N-i}` khỏi text caption / w:tblCaption.
      2) Với mỗi tbl-* được hyperlink/PAGEREF tham chiếu mà chưa có bookmark,
         gắn bookmark vào đoạn caption "Bảng N.i. ..." tương ứng.
    Chạy TRƯỚC format_figure_list() để PAGEREF vừa chèn đã trỏ đúng bookmark."""
    from docx import Document
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    doc = Document(str(docx_path))
    existing = {
        el.get(qn("w:name"))
        for el in doc.element.body.iter(qn("w:bookmarkStart"))
        if el.get(qn("w:name"))
    }

    # 1) Strip literal {#tbl-...} còn sót trong caption text + tblCaption attr.
    stripped = 0
    recovered_from_literal: dict[str, object] = {}
    for p in doc.paragraphs:
        full = "".join(t.text or "" for t in p._p.iter(qn("w:t")))
        m = LITERAL_TBL_ID_RE.search(full)
        if not m:
            continue
        anchor = m.group(1)
        for t in p._p.iter(qn("w:t")):
            if t.text and "{#tbl-" in t.text:
                t.text = LITERAL_TBL_ID_RE.sub("", t.text)
        stripped += 1
        recovered_from_literal[anchor] = p

    for el in doc.element.body.iter(qn("w:tblCaption")):
        val = el.get(qn("w:val")) or ""
        if LITERAL_TBL_ID_RE.search(val):
            el.set(qn("w:val"), LITERAL_TBL_ID_RE.sub("", val).strip())
            stripped += 1

    # 2) Thu thập tbl-* được tham chiếu (hyperlink trong danh mục).
    needed: set[str] = set()
    for hl in doc.element.body.iter(qn("w:hyperlink")):
        anchor = hl.get(qn("w:anchor"))
        if anchor and TBL_ANCHOR_RE.match(anchor):
            needed.add(anchor)
    needed.update(recovered_from_literal)

    # Map "Bảng N.i." -> paragraph caption (bỏ qua dòng trong DANH MỤC).
    caption_by_anchor: dict[str, object] = {}
    in_list = False
    for p in doc.paragraphs:
        style = p.style.name if p.style is not None else ""
        text = p.text.strip()
        if style.startswith("Heading") and text == "DANH MỤC BẢNG BIỂU":
            in_list = True
            continue
        if in_list and style.startswith("Heading"):
            in_list = False
        if in_list:
            continue
        m = BANG_CAPTION_NUM_RE.match(text)
        if m:
            caption_by_anchor[f"tbl-{m.group(1)}-{m.group(2)}"] = p

    next_id = _next_bookmark_id(doc)
    created = 0
    missing = []
    for anchor in sorted(needed):
        if anchor in existing:
            continue
        target = recovered_from_literal.get(anchor) or caption_by_anchor.get(anchor)
        if target is None:
            missing.append(anchor)
            continue
        _ensure_bookmark_on_paragraph(target, anchor, next_id, OxmlElement, qn)
        existing.add(anchor)
        next_id += 1
        created += 1

    doc.save(str(docx_path))
    parts = []
    if stripped:
        parts.append(f"gỡ {stripped} literal {{#tbl-...}} khỏi caption")
    if created:
        parts.append(f"tạo {created} bookmark tbl-* còn thiếu")
    if missing:
        parts.append(f"không tìm thấy caption cho {len(missing)} anchor ({', '.join(missing[:5])}"
                     f"{'...' if len(missing) > 5 else ''})")
    if parts:
        print("  - Đã đảm bảo bookmark bảng: " + "; ".join(parts) + ".")
    else:
        print("  - Bookmark bảng tbl-* đã đủ, không cần sửa thêm.")


def format_figure_list(docx_path: Path) -> None:
    """Định dạng lại các mục "DANH MỤC HÌNH ẢNH" / "DANH MỤC BẢNG BIỂU"
    thành kiểu "Table of Figures"/"Table of Tables" chuẩn của Word: mỗi
    dòng 1 hình/bảng, chữ đen (không còn màu xanh/gạch chân mặc định của
    hyperlink - vẫn click nhảy tới đúng vị trí được như thường), có
    tab-leader chấm (......) kéo tới lề phải trang + field PAGEREF hiện số
    trang thật (tự cập nhật khi bấm Update Field) - không còn khung bảng.

    refer_scripts/merge_bao_cao.py đã tạo sẵn mỗi dòng trong các mục này là 1
    đoạn văn CHỈ chứa đúng 1 hyperlink nội bộ [Nhãn](#fig-N-i) hoặc
    [Nhãn](#tbl-N-i) - hàm này tìm các đoạn văn đó (nằm giữa heading mục và
    heading kế tiếp) rồi thêm tab-stop + field vào cuối."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    doc = Document(str(docx_path))
    section = doc.sections[0]
    usable_width = section.page_width - section.left_margin - section.right_margin

    def _is_heading(p) -> bool:
        return p.style is not None and p.style.name.startswith("Heading")

    counts = {heading: 0 for heading, _ in LIST_SECTIONS}
    current_heading = None
    for p in doc.paragraphs:
        if current_heading is None:
            if _is_heading(p) and p.text.strip() in counts:
                current_heading = p.text.strip()
            continue
        if _is_heading(p):
            current_heading = p.text.strip() if p.text.strip() in counts else None
            continue

        anchor_re = dict(LIST_SECTIONS)[current_heading]
        hyperlinks = p._p.findall(qn("w:hyperlink"))
        if len(hyperlinks) != 1:
            continue
        anchor = hyperlinks[0].get(qn("w:anchor"))
        if not anchor or not anchor_re.match(anchor):
            continue

        _strip_hyperlink_color(hyperlinks[0], OxmlElement, qn)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.tab_stops.add_tab_stop(
            usable_width, WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS
        )
        p.add_run("\t")
        _add_pageref_field(p, anchor, OxmlElement, qn)
        counts[current_heading] += 1

    doc.save(str(docx_path))
    for heading, count in counts.items():
        print(f"  - Đã định dạng {count} dòng trong {heading} (tab-leader chấm + số trang).")


def add_page_numbering(docx_path: Path) -> None:
    """Thêm footer chỉ chứa số trang, canh giữa, có khung kẻ đôi phía trên
    (khớp mẫu báo cáo chính thức - không có chữ "Trang" và không canh phải)
    cho toàn bộ nội dung từ heading '# Chương 1' trở đi, đánh số lại từ
    trang 1 tại đó - phần trước đó (Mục lục, Danh mục hình ảnh, Danh mục
    bảng biểu) giữ nguyên footer cũ (chỉ có khung kẻ đôi, không có số
    trang), không được đánh số.

    Kỹ thuật: dùng `Document.add_section()` của python-docx để tạo đúng 1
    section + 1 footer part MỚI - tự động đăng ký relationship và
    content-type override cho part mới trong package, không cần tự tay
    thêm (điều mà thao tác lxml thuần không tiện làm). python-docx luôn
    thêm section mới ở CUỐI body, nên sau khi tạo xong, hàm này DỜI đoạn
    văn mang section-break vừa tạo (đánh dấu kết thúc section CŨ, v.d giữ
    nguyên footer cũ) tới đúng vị trí ngay trước heading "Chương 1"; sectPr
    MỚI vẫn ở cuối body như bình thường, trở thành section bao trùm từ
    Chương 1 tới hết báo cáo, và được gắn thêm `<w:pgNumType w:start="1"/>`
    để trang đầu tiên của section này tính là trang 1.

    Chạy hàm này TRƯỚC prepend_cover() trong convert() - tại thời điểm này
    docx chỉ có đúng 1 section (chưa có trang bìa), nên "cuối body" cũng
    chính là sau đoạn cuối cùng của toàn báo cáo, đơn giản hoá việc dời vị
    trí (không phải tính lại offset qua nhiều section)."""
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    doc = Document(str(docx_path))
    body = doc.element.body

    chapter1_p = next(
        (p for p in doc.paragraphs
         if p.style is not None and p.style.name == "Heading 1"
         and p.text.strip().startswith("Chương 1")),
        None,
    )
    if chapter1_p is None:
        print("  [!] Không tìm thấy heading 'Chương 1' - bỏ qua đánh số trang.",
              file=sys.stderr)
        return

    new_section = doc.add_section(WD_SECTION.NEW_PAGE)
    new_section.footer.is_linked_to_previous = False
    footer_p = new_section.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_pPr = footer_p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    top_border = OxmlElement("w:top")
    top_border.set(qn("w:val"), "double")
    top_border.set(qn("w:sz"), "6")
    top_border.set(qn("w:space"), "1")
    top_border.set(qn("w:color"), "auto")
    pBdr.append(top_border)
    footer_pPr.append(pBdr)
    _add_field(footer_p, " PAGE ", "1", OxmlElement, qn)

    # add_section() vừa thêm 2 phần tử ở cuối body: 1 đoạn văn rỗng chứa
    # sectPr-break (đóng section CŨ) rồi tới sectPr MỚI (cuối cùng của
    # body) - dời đoạn văn đó tới ngay trước heading "Chương 1".
    break_para = body[-2]
    break_pPr = break_para.find(qn("w:pPr")) if break_para.tag == qn("w:p") else None
    if break_pPr is None or break_pPr.find(qn("w:sectPr")) is None:
        print("  [!] Cấu trúc add_section() không như mong đợi - bỏ qua đánh số trang.",
              file=sys.stderr)
        return
    body.remove(break_para)
    chapter1_p._p.addprevious(break_para)

    # Đánh số lại trang từ 1 cho section mới (Chương 1 -> hết báo cáo).
    # pgNumType đứng sau pgMar trong CT_SectPrBase - sectPr hiện tại chỉ có
    # footerReference/pgSz/pgMar (không có phần tử nào đứng sau pgNumType
    # theo schema), nên append() thẳng ở cuối là hợp lệ.
    final_sectPr = body[-1]
    pgNumType = OxmlElement("w:pgNumType")
    pgNumType.set(qn("w:start"), "1")
    final_sectPr.append(pgNumType)

    doc.save(str(docx_path))
    print("  - Đã thêm footer số trang (canh giữa, khung kẻ đôi, khớp mẫu chính thức) "
          "+ đánh số lại từ trang 1, bắt đầu từ Chương 1.")


def _strip_header_footer_refs(sectPr, qn) -> None:
    """Gỡ headerReference/footerReference khỏi 1 sectPr copy từ bia.docx -
    các reference này trỏ tới r:id trong _rels/document.xml.rels của
    bia.docx (header/footer rỗng), không tồn tại trong file .docx đích nên
    nếu giữ nguyên sẽ tạo r:id treo (dangling), làm hỏng file khi Word mở."""
    for tag in ("w:headerReference", "w:footerReference"):
        for el in sectPr.findall(qn(tag)):
            sectPr.remove(el)


def _paragraph_text(el, qn) -> str:
    return "".join(t.text or "" for t in el.iter(qn("w:t"))).strip()


def _paragraph_style_id(el, qn) -> str | None:
    pPr = el.find(qn("w:pPr"))
    if pPr is None:
        return None
    pStyle = pPr.find(qn("w:pStyle"))
    if pStyle is None:
        return None
    return pStyle.get(qn("w:val"))


def _paragraph_has_sectpr(el, qn) -> bool:
    pPr = el.find(qn("w:pPr"))
    return pPr is not None and pPr.find(qn("w:sectPr")) is not None


def _strip_acknowledgements_from_cover(content_els: list, qn) -> tuple[list, bool]:
    """Gỡ trang Lời cảm ơn khỏi phần tử cover (bia.docx từng chứa trang này).

    Nội dung lời cảm ơn lấy từ loi-cam-on.md qua merge + place_acknowledgements_before_toc().
    Trả về (els_mới, đã_gỡ).
    """
    start = None
    for i, el in enumerate(content_els):
        if el.tag != qn("w:p"):
            continue
        if _paragraph_text(el, qn) == "LỜI CẢM ƠN":
            start = i
            break
    if start is None:
        return content_els, False
    return content_els[:start], True


def place_acknowledgements_before_toc(docx_path: Path) -> None:
    """Đặt khối Lời cảm ơn (từ loi-cam-on.md, pandoc biến thành Heading 1)
    ngay trước Mục lục: đổi title sang style Title (canh giữa, đậm, không
    vào outline như Heading 1), gỡ dòng LỜI CẢM ƠN khỏi field TOC nếu có,
    rồi dời cả khối lên trước SDT/TOCHeading.

    Chạy SAU pandoc --toc, TRƯỘC prepend_cover() — thứ tự cuối cùng:
    bìa + phụ bìa (bia.docx) → Lời cảm ơn → MỤC LỤC → danh mục / chương.
    """
    from docx import Document
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    doc = Document(str(docx_path))
    body = doc.element.body
    children = list(body)

    ack_start = None
    for i, el in enumerate(children):
        if el.tag != qn("w:p"):
            continue
        if _paragraph_text(el, qn) != "LỜI CẢM ƠN":
            continue
        style = _paragraph_style_id(el, qn)
        if style in (None, "Heading1", "Title", "TOCHeading"):
            ack_start = i
            break
    if ack_start is None:
        print("  [!] Không tìm thấy heading LỜI CẢM ƠN trong body, bỏ qua sắp xếp.")
        return

    ack_els = [children[ack_start]]
    for j in range(ack_start + 1, len(children)):
        el = children[j]
        if el.tag == qn("w:sdt"):
            break
        if el.tag == qn("w:p"):
            style = _paragraph_style_id(el, qn)
            text = _paragraph_text(el, qn)
            if style == "TOCHeading" or text == "MỤC LỤC":
                break
            if style == "Heading1":
                break
        ack_els.append(el)

    # Title: canh giữa, đậm (như trang Lời cảm ơn trong mẫu), không pageBreakBefore của Heading1
    title_p = ack_els[0]
    pPr = title_p.find(qn("w:pPr"))
    if pPr is None:
        pPr = OxmlElement("w:pPr")
        title_p.insert(0, pPr)
    pStyle = pPr.find(qn("w:pStyle"))
    if pStyle is None:
        pStyle = OxmlElement("w:pStyle")
        pPr.insert(0, pStyle)
    pStyle.set(qn("w:val"), "Title")
    jc = pPr.find(qn("w:jc"))
    if jc is None:
        jc = OxmlElement("w:jc")
        pPr.append(jc)
    jc.set(qn("w:val"), "center")

    # Gỡ entry LỜI CẢM ƠN khỏi field TOC (nếu pandoc đã đưa Heading 1 vào mục lục)
    removed_toc = 0
    for el in list(body):
        if el.tag != qn("w:sdt"):
            continue
        for p in list(el.iter(qn("w:p"))):
            text = _paragraph_text(p, qn)
            style = _paragraph_style_id(p, qn)
            if style == "TOCHeading":
                continue
            if text == "LỜI CẢM ƠN" or text.startswith("LỜI CẢM ƠN"):
                parent = p.getparent()
                if parent is not None:
                    parent.remove(p)
                    removed_toc += 1

    # Điểm chèn: ngay trước TOC (SDT chứa mục lục, hoặc đoạn TOCHeading)
    toc_anchor = None
    for el in body:
        if el.tag == qn("w:sdt"):
            # SDT đầu tiên thường là TOC
            texts = "".join(t.text or "" for t in el.iter(qn("w:t")))
            if "MỤC LỤC" in texts or "DANH MỤC" in texts or el is body[0]:
                toc_anchor = el
                break
        if el.tag == qn("w:p") and (
            _paragraph_style_id(el, qn) == "TOCHeading"
            or _paragraph_text(el, qn) == "MỤC LỤC"
        ):
            toc_anchor = el
            break
    if toc_anchor is None:
        toc_anchor = body[0] if len(body) else None

    for el in ack_els:
        body.remove(el)
    if toc_anchor is not None and toc_anchor in list(body):
        for el in ack_els:
            toc_anchor.addprevious(el)
    else:
        for el in reversed(ack_els):
            body.insert(0, el)

    doc.save(str(docx_path))
    extra = f", gỡ {removed_toc} dòng khỏi TOC" if removed_toc else ""
    print(f"  - Đã đặt Lời cảm ơn trước Mục lục (style Title{extra}).")


def prepend_cover(docx_path: Path, cover_path: Path) -> None:
    """Chèn nội dung trang bìa từ `cover_path` (refer_scripts/bia.docx - trang bìa
    ngoài + trang phụ bìa, mỗi trang là 1 section riêng với lề/khung viền
    khác nhau) vào ĐẦU `docx_path` (đã convert xong từ markdown), thay vì
    render trang bìa bằng markdown/pandoc (không tái tạo được khung viền +
    layout 2 trang bìa chuẩn).

    Trang Lời cảm ơn trong bia.docx (nếu còn) bị GỠ khi chèn — nội dung lấy
    từ loi-cam-on.md đã được place_acknowledgements_before_toc() xếp trước
    Mục lục.

    Giữ nguyên section break giữa 2 trang bìa/phụ bìa, rồi chèn 1 section
    break sau phụ bìa để nội dung báo cáo gốc tiếp tục dùng đúng section
    (lề/header/footer) của chính báo cáo, không bị đổi theo lề trang bìa."""
    from copy import deepcopy
    from docx import Document
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    cover = Document(str(cover_path))
    target = Document(str(docx_path))

    cover_children = list(cover.element.body)
    if not cover_children or cover_children[-1].tag != qn("w:sectPr"):
        print(f"  [!] {cover_path.name}: cấu trúc không như mong đợi (thiếu sectPr cuối), "
              "bỏ qua chèn trang bìa.", file=sys.stderr)
        return

    content_els = [deepcopy(el) for el in cover_children[:-1]]
    trailing_sectpr = deepcopy(cover_children[-1])
    _strip_header_footer_refs(trailing_sectpr, qn)
    for el in content_els:
        for sect in el.findall(f".//{qn('w:sectPr')}"):
            _strip_header_footer_refs(sect, qn)

    content_els, stripped_ack = _strip_acknowledgements_from_cover(content_els, qn)

    ends_with_sect = bool(content_els) and content_els[-1].tag == qn("w:p") and (
        _paragraph_has_sectpr(content_els[-1], qn)
    )
    if ends_with_sect:
        # Phụ bìa đã kết thúc bằng sectBreak — không thêm break_para nữa
        # (tránh tạo section trống giữa phụ bìa và Lời cảm ơn).
        last_pPr = content_els[-1].find(qn("w:pPr"))
        if last_pPr is not None:
            for sect in last_pPr.findall(qn("w:sectPr")):
                _strip_header_footer_refs(sect, qn)
    else:
        break_para = OxmlElement("w:p")
        pPr = OxmlElement("w:pPr")
        pPr.append(trailing_sectpr)
        break_para.append(pPr)
        content_els.append(break_para)

    target_body = target.element.body
    anchor = target_body[0] if len(target_body) else None
    for el in content_els:
        if anchor is not None:
            anchor.addprevious(el)
        else:
            target_body.append(el)

    target.save(str(docx_path))
    note = ", đã bỏ trang Lời cảm ơn cũ trong bìa" if stripped_ack else ""
    print(f"  - Đã chèn trang bìa từ {cover_path.name} vào đầu tài liệu{note}.")


def check_pandoc() -> None:
    if shutil.which("pandoc") is None:
        sys.exit(
            "Không tìm thấy pandoc trong PATH.\n"
            "Cài đặt: https://pandoc.org/installing.html\n"
            "  - macOS:   brew install pandoc\n"
            "  - Ubuntu:  sudo apt install pandoc\n"
            "  - Windows: choco install pandoc  (hoặc tải installer trên trang chủ)"
        )


def convert(md_path: Path, out_path: Path, reference_doc: Path | None, keep_temp: bool,
            cover_doc: Path | None = None) -> None:
    check_pandoc()
    if not md_path.is_file():
        sys.exit(f"Không tìm thấy file markdown: {md_path}")

    tmp_dir = Path(tempfile.mkdtemp(prefix="md2docx_"))
    try:
        source_text = build_docx_source(md_path, tmp_dir / "images")

        cmd = [
            "pandoc",
            # +raw_attribute: cho phép nhúng raw OOXML ngay trong markdown
            # (dạng `<w:...>`{=openxml} hoặc ```{=openxml} ... ```) nếu cần
            # về sau; hiện các field PAGEREF (mục DANH MỤC HÌNH ẢNH, xem
            # format_figure_list() dưới) được chèn bằng python-docx sau khi
            # pandoc convert xong, không cần extension này, nhưng để sẵn
            # không hại gì.
            "-f", "markdown+raw_attribute",
            "-t", "docx",
            "--resource-path", str(md_path.parent),
            # Mục lục: để pandoc tự sinh field TOC thật của Word (tự cập
            # nhật số trang khi bấm Update Field/F9), dựa theo các heading
            # Chương/mục (Heading1-3) đã có sẵn - không viết tay trong
            # markdown. Đặt ngay đầu tài liệu (trước heading đầu tiên -
            # "DANH MỤC HÌNH ẢNH", xem merge_bao_cao.py:build_report()),
            # rồi được prepend_cover() chèn trang bìa lên trước nữa.
            "--toc",
            "--toc-depth=3",
            "-M", "toc-title=MỤC LỤC",
            "-o", str(out_path),
        ]
        if reference_doc and reference_doc.is_file():
            cmd += ["--reference-doc", str(reference_doc)]
        elif reference_doc:
            print(f"  [!] Không tìm thấy reference-doc {reference_doc}, dùng style mặc định của pandoc.",
                  file=sys.stderr)

        print(f"Đang chạy pandoc -> {out_path}")
        result = subprocess.run(cmd, input=source_text.encode("utf-8"), capture_output=True)
        if result.returncode != 0:
            sys.exit(f"pandoc lỗi:\n{result.stderr.decode(errors='ignore')}")
        if result.stderr:
            print(result.stderr.decode(errors="ignore"), file=sys.stderr)

        try:
            import docx  # noqa: F401
        except ImportError:
            print(
                "  [!] Không import được python-docx — bỏ qua TOÀN BỘ bước hậu "
                "xử lý (viền bảng, canh giữa ảnh/caption, MỤC LỤC, trang bìa, "
                f"số trang, …).\n"
                f"      Python đang dùng: {sys.executable}\n"
                "      Cài vào đúng môi trường này, ví dụ:\n"
                f"        {sys.executable} -m pip install python-docx\n"
                "      (Nếu đang bật .venv của repo: pip install python-docx)",
                file=sys.stderr,
            )
        else:
            set_body_alignment_justify(out_path)
            flatten_inline_code_style(out_path)
            flatten_markdown_emphasis(out_path)
            fix_toc_heading_style(out_path)
            fix_table_widths(out_path)
            fix_image_widths(out_path)
            fix_list_alignment(out_path)
            center_images(out_path)
            ensure_table_bookmarks(out_path)
            format_figure_list(out_path)
            place_acknowledgements_before_toc(out_path)
            add_page_numbering(out_path)
            if cover_doc and cover_doc.is_file():
                prepend_cover(out_path, cover_doc)
            elif cover_doc:
                print(f"  [!] Không tìm thấy cover-doc {cover_doc}, bỏ qua chèn trang bìa.",
                      file=sys.stderr)

        print(f"Xong: {out_path}")
    finally:
        if keep_temp:
            print(f"(Giữ lại thư mục ảnh tạm để debug: {tmp_dir})")
        else:
            shutil.rmtree(tmp_dir, ignore_errors=True)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.split("Cách dùng")[0],
                                      formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("md_file", nargs="?", default=str(DEFAULT_MD),
                         help=f"File markdown cần convert (mặc định: {DEFAULT_MD.relative_to(PROJECT_ROOT)})")
    parser.add_argument("-o", "--output", default=None,
                         help="File .docx đầu ra (mặc định: cùng tên, cùng thư mục file md)")
    parser.add_argument("--reference-doc", default=str(DEFAULT_REFERENCE_DOC),
                         help="File .docx dùng làm style/template cho pandoc (--reference-doc)")
    parser.add_argument("--cover-doc", default=str(DEFAULT_COVER_DOC),
                         help="File .docx dùng làm trang bìa, chèn vào đầu tài liệu sau khi "
                              f"convert (mặc định: {DEFAULT_COVER_DOC.relative_to(PROJECT_ROOT)})")
    parser.add_argument("--no-cover", action="store_true",
                         help="Không chèn trang bìa (bỏ qua --cover-doc)")
    parser.add_argument("--keep-temp", action="store_true",
                         help="Giữ lại thư mục PNG tạm (để debug ảnh SVG convert lỗi)")
    args = parser.parse_args()

    md_path = Path(args.md_file).resolve()
    out_path = Path(args.output).resolve() if args.output else md_path.with_suffix(".docx")
    reference_doc = Path(args.reference_doc).resolve() if args.reference_doc else None
    cover_doc = None if args.no_cover else (Path(args.cover_doc).resolve() if args.cover_doc else None)

    convert(md_path, out_path, reference_doc, args.keep_temp, cover_doc)


if __name__ == "__main__":
    main()
